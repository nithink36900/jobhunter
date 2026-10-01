"""Resume parsing, job fetching (Apify), scoring and SQLite tracking."""
import json
import os
import re
import sqlite3
import time
from datetime import datetime

SKILL_VOCAB = [
    "wireshark", "nmap", "metasploit", "burp suite", "sqlmap", "steghide", "python", "java",
    "kali", "linux", "ubuntu", "windows", "vapt", "penetration testing", "pentest", "ethical hacking",
    "vulnerability assessment", "vulnerability management", "threat detection", "threat intelligence",
    "network security", "firewall", "vpn", "dns", "tcp/ip", "encryption", "cryptography", "steganography",
    "siem", "splunk", "soc", "incident response", "owasp", "ceh", "blockchain", "exploit", "malware",
    "forensic", "iam", "access control", "nessus", "qradar", "sentinel", "edr", "grc", "iso 27001",
    "web application security", "appsec", "automation", "scripting", "ctf", "cloud security", "aws",
]
ROLE_WORDS = ["security", "cyber", "soc", "vapt", "pentest", "penetration", "infosec", "threat",
              "vulnerability", "appsec", "red team", "blue team", "forensic", "iam", "grc"]


# ------------------------------ config ------------------------------
def load_cfg(path="config.yaml", settings="data/settings.json"):
    """config.yaml, with anything saved from the web Settings page laid on top."""
    import yaml
    cfg = yaml.safe_load(open(path, encoding="utf-8"))
    if os.path.exists(settings):
        try:
            ov = json.load(open(settings, encoding="utf-8"))
        except ValueError:
            ov = {}  # damaged settings file: fall back to config.yaml instead of crashing
        cfg["search"].update({k: ov[k] for k in ("titles", "locations", "posted_within", "experience_levels") if k in ov})
        cfg["matching"].update({k: ov[k] for k in ("min_score", "top_n") if k in ov})
    return cfg


# ------------------------------ resume ------------------------------
def read_resume(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    return open(path, encoding="utf-8", errors="ignore").read()


def resume_skills(text):
    low = text.lower()
    return sorted({s for s in SKILL_VOCAB if re.search(r"(?<![a-z])" + re.escape(s) + r"(?![a-z])", low)})


# ------------------------------ fetching ------------------------------
def _fill(template, values):
    """Replace '{key}' strings with real values (keeps lists/ints as-is)."""
    if isinstance(template, dict):
        return {k: _fill(v, values) for k, v in template.items()}
    if isinstance(template, str):
        m = re.fullmatch(r"\{(\w+)\}", template)
        return values[m.group(1)] if m else template
    return template


def normalise(item):
    g = lambda *ks: next((item[k] for k in ks if item.get(k)), "")
    return {
        "title": g("title", "jobTitle", "positionName", "name"),
        "company": g("companyName", "company", "companyname"),
        "location": g("location", "jobLocation", "city"),
        "posted": g("postedTime", "postedAt", "postingDateParsed", "publishedAt", "datePosted"),
        "job_url": g("jobUrl", "url", "link", "jobLink"),
        "apply_url": g("applyUrl", "externalApplyLink", "applyLink", "jobUrl", "url", "link"),
        "level": g("experienceLevel", "seniority", "experience"),
        "type": g("contractType", "employmentType", "jobType"),
        "work_type": g("workType", "workplaceType"),
        "applicants": g("applicationsCount", "applicants"),
        "description": g("description", "descriptionText", "jobDescription"),
    }


def get_apify_token(token=None):
    """Token comes from the .env file in the project folder (saved from the web app), else the APIFY_TOKEN variable."""
    token = (token or "").strip()
    if not token and os.path.exists(".env"):
        for line in open(".env", encoding="utf-8"):
            key, _, val = line.strip().partition("=")
            if key.strip() == "APIFY_TOKEN":
                token = val.strip().strip("\"'")
    if not token:
        token = os.environ.get("APIFY_TOKEN", "").strip()
    if not token:
        raise SystemExit(
            "Apify token not found.\n"
            "Open the web app, go to Settings and paste your token there.\n"
            "Get a token at console.apify.com > Settings > API & integrations.")
    return token


def token_present():
    try:
        get_apify_token()
        return True
    except SystemExit:
        return False


def save_token(token, path=".env"):
    """Store the token in .env (keeps any other lines)."""
    keep = []
    if os.path.exists(path):
        keep = [l.rstrip("\n") for l in open(path, encoding="utf-8") if not l.startswith("APIFY_TOKEN")]
    with open(path + ".tmp", "w", encoding="utf-8") as f:
        f.write("\n".join(keep + ["APIFY_TOKEN=" + token]) + "\n")
    os.replace(path + ".tmp", path)


def _dataset_id(run):
    """Old apify-client returns a dict, new versions return a Run object. Handle both."""
    if run is None:
        raise RuntimeError("Apify did not start the job. Check your Apify account and credits at console.apify.com.")
    if isinstance(run, dict):
        return run["defaultDatasetId"]
    for name in ("default_dataset_id", "defaultDatasetId"):
        if getattr(run, name, None):
            return getattr(run, name)
    if hasattr(run, "model_dump"):
        d = run.model_dump(by_alias=True)
        if d.get("defaultDatasetId") or d.get("default_dataset_id"):
            return d.get("defaultDatasetId") or d.get("default_dataset_id")
    raise RuntimeError("Could not read the results id from Apify's reply.")


def fetch_jobs(cfg, token=None, log=print):
    token = get_apify_token(token)  # check the token first, so a missing one gives a clear message
    from apify_client import ApifyClient
    client = ApifyClient(token)
    s = cfg["search"]
    values = {"titles": s["titles"], "locations": s["locations"], "rows": s["rows_per_search"],
              "posted_within": s["posted_within"], "experience_levels": s["experience_levels"]}
    jobs = []
    for src in cfg["sources"]:
        if not src.get("enabled"):
            continue
        log(f"[fetch] {src['name']} via {src['actor']} ...")
        run = client.actor(src["actor"]).call(run_input=_fill(src["input_template"], values))
        for it in client.dataset(_dataset_id(run)).iterate_items():
            j = normalise(it)
            j["source"] = src["name"]
            if j["title"] and j["apply_url"]:
                jobs.append(j)
        log(f"[fetch] {src['name']}: {len(jobs)} total so far")
    return jobs


# ------------------------------ scoring ------------------------------
def _age_hours(posted):
    p = str(posted).lower()
    m = re.search(r"(\d+)\s*(minute|hour|day|week|month)", p)
    if not m:
        return 72
    n, unit = int(m.group(1)), m.group(2)
    return n * {"minute": 1 / 60, "hour": 1, "day": 24, "week": 168, "month": 720}[unit]


def _applicants(txt):
    m = re.search(r"(\d+)", str(txt))
    if "over" in str(txt).lower() and m:
        return int(m.group(1)) + 1
    return int(m.group(1)) if m else 100


def score_job(job, cfg, skills):
    w, mt = cfg["matching"]["weights"], cfg["matching"]
    title, desc = job["title"].lower(), job["description"].lower()
    loc = f"{job['location']} {job['work_type']}".lower()

    # title: role keywords + closeness to searched titles; excluded words kill the score
    entry_title = re.search(r"trainee|intern|junior|fresher|entry|graduate|\bassociate\b", title)
    if not entry_title and any(re.search(rf"\b{re.escape(x)}\b", title) for x in mt["exclude_title_words"]):
        return 0, "senior title"
    hits = sum(k in title for k in ROLE_WORDS)
    wanted = max((len(set(title.split()) & set(t.lower().split())) / len(t.split())
                  for t in cfg["search"]["titles"]), default=0)
    t = w["title"] * min(1.0, 0.5 * min(hits, 2) + 0.5 * wanted) if hits else 0

    # skills overlap with description (falls back to neutral if no description)
    if desc:
        found = [s for s in skills if s in desc]
        sk = w["skills"] * min(1.0, len(found) / 5)
    else:
        sk = w["skills"] * 0.4

    # level
    lvl = f"{job['level']} {title}".lower()
    if any(k in lvl for k in ("intern", "entry", "junior", "trainee", "associate", "fresher", "graduate")):
        lv = w["level"]
    elif "not applicable" in lvl or not job["level"]:
        lv = w["level"] * 0.5
    else:
        lv = w["level"] * 0.2
    yrs = [int(x) for x in re.findall(r"(\d+)\+?\s*(?:-\s*\d+\s*)?years", desc)]
    if yrs and min(yrs) > mt["max_years_experience"]:
        lv *= 0.3

    # location
    pref = cfg["search"]["preferred_locations"]
    lo = w["location"] if any(p in loc for p in pref) else w["location"] * (0.5 if "india" in loc else 0.1)

    # recency + competition
    age = _age_hours(job["posted"])
    rc = w["recency"] * (1 if age <= 24 else 0.8 if age <= 72 else 0.5 if age <= 168 else 0.2)
    ap = _applicants(job["applicants"])
    cp = w["competition"] * (1 if ap < 25 else 0.7 if ap < 75 else 0.4 if ap < 150 else 0.2)

    total = round(t + sk + lv + lo + rc + cp)
    return min(total, 99), ""


def rank(jobs, cfg, skills):
    for j in jobs:
        j["score"], j["note"] = score_job(j, cfg, skills)
    seen, out = set(), []
    for j in sorted(jobs, key=lambda x: -x["score"]):
        key = (j["title"].lower(), j["company"].lower(), j["location"].lower())
        if key in seen or j["score"] < cfg["matching"]["min_score"]:
            continue
        seen.add(key)
        out.append(j)
    return out[: cfg["matching"]["top_n"]]


# ------------------------------ tracking ------------------------------
def ensure_schema(db):
    db.execute("""CREATE TABLE IF NOT EXISTS jobs(
        apply_url TEXT PRIMARY KEY, title TEXT, company TEXT, score INT,
        status TEXT DEFAULT 'new', data TEXT, updated TEXT)""")
    if "notes" not in [r[1] for r in db.execute("PRAGMA table_info(jobs)")]:
        db.execute("ALTER TABLE jobs ADD COLUMN notes TEXT DEFAULT ''")
    db.commit()


class Tracker:
    def __init__(self, path="data/jobs.db"):
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        self.db = sqlite3.connect(path)
        ensure_schema(self.db)

    def save(self, jobs):
        new = 0
        for j in jobs:
            cur = self.db.execute("SELECT 1 FROM jobs WHERE apply_url=?", (j["apply_url"],)).fetchone()
            if cur:
                self.db.execute("UPDATE jobs SET score=? WHERE apply_url=?", (j["score"], j["apply_url"]))
            else:
                new += 1
                self.db.execute("INSERT INTO jobs(apply_url,title,company,score,status,data,updated) VALUES(?,?,?,?,?,?,?)",
                                (j["apply_url"], j["title"], j["company"], j["score"], "new",
                                 json.dumps(j), datetime.now().isoformat()))
        self.db.commit()
        return new

    def pending(self, limit):
        rows = self.db.execute("SELECT data FROM jobs WHERE status='new' ORDER BY score DESC LIMIT ?",
                               (limit,)).fetchall()
        return [json.loads(r[0]) for r in rows]

    def mark(self, url, status):
        self.db.execute("UPDATE jobs SET status=?, updated=? WHERE apply_url=?",
                        (status, datetime.now().isoformat(), url))
        self.db.commit()

    def all(self):
        return self.db.execute(
            "SELECT score,title,company,status,apply_url FROM jobs ORDER BY score DESC").fetchall()


def markdown_table(jobs):
    lines = ["| Apply Link | Rank | Match Score | Role | Company |", "|---|---|---|---|---|"]
    for i, j in enumerate(jobs, 1):
        lines.append(f"| [Apply]({j['apply_url']}) | {i} | {j['score']}% | {j['title']} "
                     f"({j['location']}) | {j['company']} |")
    return "\n".join(lines)
